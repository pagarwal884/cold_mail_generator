import os
import requests
import json
import re
import google.generativeai as genai
from django.conf import settings
from dotenv import load_dotenv
import magic
import tempfile
from pdfminer.high_level import extract_text
from docx import Document

load_dotenv()

def generate_cold_mail_gemini_api(resume_content, target_company, role_applied_for, tone):
    api_key = os.getenv('GEMINI_API_KEY')
    if not api_key:
        raise Exception("Gemini API Key is not configured. Please set GEMINI_API_KEY in your .env file.")

    prompt = f'''
You are an expert cold email writer. Your task is to generate a professional and engaging cold email, including a compelling subject line.

**Context:**
The user is applying for the role of **{role_applied_for}** (or similar roles) at **{target_company}**.
The user wants the email to have a {tone} tone.

**User's Resume Highlights / Key Information:**
---
{resume_content}
---

**Instructions:**
1. Craft a concise and compelling cold email body.
2. The email should be suitable for sending to a hiring manager or recruiter at {target_company} for the role of {role_applied_for} or related positions.
3. Highlight relevant skills and experiences from the resume information provided that specifically align with the **{role_applied_for}** role and potential needs at {target_company}.
4. Express strong interest in the **{role_applied_for}** role and in {target_company}.
5. Include a clear call to action (e.g., requesting a brief informational interview, discussion about suitability for the {role_applied_for} role, or expressing eagerness to learn about relevant openings).
6. Maintain the specified {tone} throughout the email.
7. **The output MUST be a valid JSON object with two string keys: "subject" and "body".**
   - "subject": A compelling and concise subject line for this email (around 5-10 words).
   - "body": The full email content.
8. **For the "body":**
   - Start directly with an appropriate salutation (e.g., "Dear Hiring Manager,").
   - End with a professional closing phrase (e.g., "Sincerely,", "Best regards,"). **Do NOT add any name or signature line after this closing phrase.** The user will add their own.
9. Do not include placeholders like "[Your Name]", "[Your Contact Information]", "[Date]" *anywhere* in the email subject or body.
10. Ensure the email body is well-structured, grammatically correct, and professional.
'''

    genai.configure(api_key=api_key)
    model = genai.GenerativeModel('gemini-2.0-flash')
    try:
        response = model.generate_content(prompt)
        raw_text = response.text
        fence_regex = r'^```(?:json)?\s*\n?(.*?)\n?\s*```$'
        match = re.match(fence_regex, raw_text, re.DOTALL)
        if match:
            json_str = match.group(1).strip()
        else:
            json_str = raw_text.strip()
        parsed = json.loads(json_str)
        subject = parsed.get('subject', '').strip()
        body = parsed.get('body', '').strip()
        if not subject or not body:
            raise Exception("Generated email response from AI is missing the subject or body.")
        return subject, body
    except Exception as e:
        raise Exception(f"Failed to generate or parse Gemini AI response: {str(e)}")

def parse_resume(file):
    # Ensure file pointer is at start
    if hasattr(file, 'seek') and callable(file.seek):
        file.seek(0)
    
    # Read initial bytes for type detection
    file_sample = file.read(1024)
    file.seek(0)  # Reset file pointer
    
    # Detect MIME type using magic
    try:
        file_type = magic.from_buffer(file_sample, mime=True)
    except AttributeError:
        # Fallback if magic doesn't have from_buffer
        import mimetypes
        file_type = mimetypes.guess_type(file.name)[0]
    
    text = ""

    if file_type == 'application/pdf':
        with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
            for chunk in file.chunks() if hasattr(file, 'chunks') else [file.read()]:
                tmp.write(chunk)
            tmp_path = tmp.name
        try:
            text = extract_text(tmp_path)
        finally:
            os.unlink(tmp_path)
    
    elif file_type in [
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'application/msword'
    ]:
        try:
            doc = Document(file)
            text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        except Exception as e:
            raise Exception(f"Error parsing Word document: {str(e)}")
    
    else:
        raise Exception(f"Unsupported file type: {file_type}")
    
    return text